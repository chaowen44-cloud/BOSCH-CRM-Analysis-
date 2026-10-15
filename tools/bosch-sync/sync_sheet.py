"""Parse the BOSCH Retail Membership Google Sheet (xlsx export) into one JSON doc per month.

Usage: python3 sync_sheet.py <drive-tool-result.json | file.xlsx> <out_dir>

Each tab named YYYY-MM or YYYYMM (e.g. 2026-09) is one month; other tabs (copies etc.) are ignored.
The dashed form is what the page's live read needs; the output id is always YYYYMM.
Per tab: row 3 = DDS, row 4 = Showroom (cols B/C/D = cumulative / monthly new / closed count, kept under the key "shipped"; the optional
"轉換筆數" column, located by its header, becomes "conv");
store rows start at row 8 and run until column A is blank or a section header
("Total Retail", "1-on-1 ...") begins. Blank or non-numeric cells -> 0.
Writes <out_dir>/<YYYYMM>.json shaped {dds:{cum,monthly,shipped}, showroom:{...}, stores:[...]},
plus chat:{dds,showroom} (1-on-1 chat windows) and tags (total tags used) for months whose tab
has a "1-on-1 Chat" / "Total Tags Used" block, and chatStores:[{name,windows}, ...] for months whose
tab has a "1-on-1 Chat (by store)" block (store names in column A, chat windows in column B).
"""
import base64, json, os, re, sys
import openpyxl


def num(v):
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return 0


def conv_col(rows):
    """Index of the "轉換筆數" (conversion count) column, found from the header rows (2 and 7); None if absent."""
    for r in (rows[1], rows[6]):
        for i, v in enumerate(r):
            if isinstance(v, str) and '轉換筆數' in v:
                return i
    return None


def triple(row, cc=None):
    """cum / monthly new / closed count (column D, still keyed "shipped") and, when the tab has the column, conv."""
    t = {'cum': num(row[1]), 'monthly': num(row[2]), 'shipped': num(row[3])}
    if cc is not None and cc < len(row):
        t['conv'] = num(row[cc])
    return t


def is_chat_by_store(text):
    return text.lower().startswith('1-on-1') and 'by store' in text.lower()


def chat_block(rows):
    """The optional "1-on-1 Chat" block below the store list: DDS / Showroom chat windows in column B."""
    label = lambda r: '' if r[0] is None else str(r[0]).strip()
    start = next((i for i, r in enumerate(rows)
                  if label(r).lower().startswith('1-on-1') and not is_chat_by_store(label(r))), None)
    if start is None:
        return None
    found = {}
    for r in rows[start + 1:]:
        key = {'DDS': 'dds', 'Showroom': 'showroom'}.get(label(r))
        if key and key not in found:
            found[key] = num(r[1])
    return found if len(found) == 2 else None


def chat_stores_block(rows):
    """The optional "1-on-1 Chat (by store)" block: one row per store, chat windows in column B.
    Its header row (blank column A, "Chat Windows" in B) is skipped; the list ends at a blank row
    or the next section header."""
    label = lambda r: '' if r[0] is None else str(r[0]).strip()
    start = next((i for i, r in enumerate(rows) if is_chat_by_store(label(r))), None)
    if start is None:
        return None
    i = start + 1
    while i < len(rows) and label(rows[i]) == '':
        i += 1
    stores = []
    for r in rows[i:]:
        name = label(r)
        if name == '' or name == 'Total Retail' or name.lower().startswith(('1-on-1', 'total tags used')):
            break
        stores.append({'name': name, 'windows': num(r[1])})
    return stores or None


def tags_block(rows):
    """The optional "Total Tags Used" block: the number in column B of its "Total Retail" row."""
    label = lambda r: '' if r[0] is None else str(r[0]).strip()
    start = next((i for i, r in enumerate(rows) if label(r).lower().startswith('total tags used')), None)
    if start is None:
        return None
    for r in rows[start + 1:]:
        if label(r) == 'Total Retail' and isinstance(r[1], (int, float)):
            return num(r[1])
    return None


def load_workbook(src):
    if src.endswith('.xlsx'):
        return openpyxl.load_workbook(src, data_only=True)
    data = base64.b64decode(json.load(open(src))['content'])
    tmp = src + '.xlsx'
    open(tmp, 'wb').write(data)
    return openpyxl.load_workbook(tmp, data_only=True)


def main(src, out_dir):
    wb = load_workbook(src)
    os.makedirs(out_dir, exist_ok=True)
    for name in wb.sheetnames:
        m = re.fullmatch(r'(\d{4})-?(\d{2})', name.strip())
        if not m:
            continue
        month_id = m.group(1) + m.group(2)
        rows = list(wb[name].iter_rows(min_row=1, max_row=wb[name].max_row, values_only=True))
        cc = conv_col(rows)
        doc = {'dds': triple(rows[2], cc), 'showroom': triple(rows[3], cc), 'stores': []}
        for r in rows[7:]:
            label = '' if r[0] is None else str(r[0]).strip()
            if label == '' or label == 'Total Retail' or label.lower().startswith(('1-on-1', 'total tags used')):
                break  # end of the store list: blank row or the next section (e.g. the Total Tags / 1-on-1 Chat block)
            doc['stores'].append({'name': str(r[0]).strip(), **triple(r, cc)})
        chat = chat_block(rows)
        if chat:
            doc['chat'] = chat
        chat_stores = chat_stores_block(rows)
        if chat_stores:
            doc['chatStores'] = chat_stores
        tags = tags_block(rows)
        if tags is not None:
            doc['tags'] = tags
        json.dump(doc, open(os.path.join(out_dir, month_id + '.json'), 'w', encoding='utf-8'), ensure_ascii=False)
        print(month_id, 'DDS', doc['dds'], 'Showroom', doc['showroom'], len(doc['stores']), 'stores')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
